"""Readable Word documents using the shared saved-result presentation."""
import re
from collections import Counter
from hashlib import sha256
from pathlib import Path
from datetime import datetime, timezone
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def render_word(document, destination, progress=None):
    from core.report_document import prepare_document
    document = prepare_document(document)
    from utils.datetime_utils import presentation_dates
    document = presentation_dates(document)
    from core.report_progress import report_progress
    report_progress(progress, 'building', 'Preparing Word report evidence')
    from core.report_document import presentation_context
    context = presentation_context(document)
    labels = {'summary':'Scan Summary', 'areas_of_interest':'Areas of Interest', 'filepaths_aoi':'Path Areas of Interest',
              'filepaths':'File Inventory', 'analysis':'Advanced Analysis: Vulnerabilities, Attack Vectors & Flows',
              'rdl_review_evidence':'RDL Review Evidence', 'rdl_cross_references':'Cross References',
              'suppressed_findings':'Suppressed Findings', 'recon':'Reconnaissance',
              'source_excerpt_catalogue':'Source Excerpt Catalogue'}
    labels.update(native_reference_catalogue='Flow Reference Catalogue', review_text_catalogue='Review Text & Diagnostics')
    labels['program_graphs'] = 'Control Flow Graphs & Data Flow Graphs'
    doc = Document()
    generated = datetime.now(timezone.utc).replace(tzinfo=None)
    doc.core_properties.author = doc.core_properties.last_modified_by = 'Daksh SCRA'
    doc.core_properties.title = 'Security Analysis Report'
    doc.core_properties.subject = document.get('report', 'Full scan')
    doc.core_properties.created = doc.core_properties.modified = generated
    section = doc.sections[0]
    section.page_height, section.page_width = Inches(11.69), Inches(8.27)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Calibri', Pt(10)
    normal.font.color.rgb = RGBColor.from_string('1E2D45')
    normal.paragraph_format.space_after = Pt(5)
    for name in ('Title', 'Heading 1', 'Heading 2', 'Heading 3'):
        doc.styles[name].font.color.rgb = RGBColor.from_string('164C8C')
    header = section.header.paragraphs[0]
    header.text = 'DAKSH SCRA  |  SOURCE CODE SECURITY REVIEW'
    header.style = 'Caption'
    footer = section.footer.paragraphs[0]
    footer.text = 'Daksh SCRA  |  Security Review                                      Page '
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    doc.add_heading('Security Analysis Report', 0)
    doc.add_paragraph(document.get('report', 'Full scan'), 'Subtitle')
    from utils.report_audience import LABELS
    doc.add_paragraph('Report audience: ' + LABELS[document.get('report_audience', 'both')])
    scan = document.get('scan', {})
    for label, value in [('Project', scan.get('project')), ('Scan', scan.get('run_uuid'))]:
        if value:
            doc.add_paragraph(f'{label}: {value}')
    doc.add_paragraph('Areas of interest and static analysis evidence support human review. Recorded outcomes, assessment status, and coverage limits are preserved.')
    doc.add_paragraph('The readable evidence follows the HTML and PDF reports. Recorded evidence and the selected guidance are available in the JSON export and HTML package. The saved scan retains both guidance perspectives.')
    summary = document.get('sections', {}).get('summary', {})
    metrics = summary.get('detection_summary', {}) if isinstance(summary, dict) else {}
    table = doc.add_table(rows=0, cols=2)
    table.style = 'Light Shading Accent 1'
    for label, value in [('Files scanned', metrics.get('total_files_scanned', 'Not recorded')),
                         ('Areas of interest', metrics.get('areas_of_interest_identified', 'Not recorded')),
                         ('Path findings', metrics.get('file_paths_areas_of_interest_identified', 'Not recorded')),
                         ('Confirmed RDL findings', context['review_summary']['native_assessments']),
                         ('Grouped RDL review findings', context['review_summary']['rdl_groups']),
                         ('Inconclusive RDL checks / coverage limits', context['review_summary']['coverage_limits'])]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, str(value)
    doc.add_paragraph('Vulnerability classifications use RDL v2 outcomes. Confirmed findings mean the named RDL property matched; exploitation is not established. Detailed evidence retains classification and coverage limits.')
    doc.add_page_break()
    doc.add_heading('Contents', 1)
    toc = OxmlElement('w:fldSimple'); toc.set(qn('w:instr'), r'TOC \o "1-1" \h \z \u')
    # Cached links keep the contents useful in readers that do not update fields.
    # Word can replace this cache with its own paginated contents on update.
    def contents_link(key, title):
        if len(toc):
            run = OxmlElement('w:r'); run.append(OxmlElement('w:br')); toc.append(run)
        link = OxmlElement('w:hyperlink'); link.set(qn('w:anchor'), 'section_' + key)
        run = OxmlElement('w:r'); properties = OxmlElement('w:rPr')
        color = OxmlElement('w:color'); color.set(qn('w:val'), '164C8C'); properties.append(color); run.append(properties)
        text = OxmlElement('w:t'); text.text = title; run.append(text); link.append(run); toc.append(link)
    for key in document.get('sections', {}):
        contents_link(key, labels.get(key, key.replace('_',' ').title()))
    doc.add_paragraph()._p.append(toc)
    settings = doc.settings.element
    # Cached contents already navigates correctly. Let readers refresh it
    # explicitly instead of rebuilding a large document during initial open.
    update = OxmlElement('w:updateFields'); update.set(qn('w:val'), 'false'); settings.append(update)
    def clean(value):
        return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', lambda m: '\\u%04x' % ord(m[0]), str(value))
    body = doc.element.body
    link_targets = set()
    reference_prefixes = {'Source excerpt':'source_excerpt_', 'Flow reference':'flow_reference_',
                          'Review text':'review_text_', 'RDL finding':'rdl_finding_', 'RDL flow':'rdl_flow_'}
    reference_choices = '|'.join(re.escape(value) for value in reference_prefixes)
    reference_pattern = re.compile(r'(' + reference_choices + r') (\d+)')
    section_properties = body.sectPr
    body.remove(section_properties)
    def heading(text, level=1, bookmark=None):
        p = OxmlElement('w:p'); properties = OxmlElement('w:pPr')
        if level == 1:
            properties.append(OxmlElement('w:pageBreakBefore'))
        style = OxmlElement('w:pStyle'); style.set(qn('w:val'), f'Heading{level}')
        properties.append(style); p.append(properties)
        if bookmark:
            start = OxmlElement('w:bookmarkStart'); start.set(qn('w:id'), str(bookmark[0])); start.set(qn('w:name'), bookmark[1]); p.append(start)
        run = OxmlElement('w:r'); node = OxmlElement('w:t'); node.text = clean(text); run.append(node); p.append(run)
        if bookmark:
            end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(bookmark[0])); p.append(end)
        body.append(p)
    def paragraph(label, value, links=True):
        # Direct OOXML insertion keeps large evidence collections linear in size.
        p = OxmlElement('w:p')
        if label:
            run = OxmlElement('w:r'); prop = OxmlElement('w:rPr'); prop.append(OxmlElement('w:b')); run.append(prop)
            t = OxmlElement('w:t'); t.text = clean(label) + ': '; t.set(qn('xml:space'), 'preserve'); run.append(t); p.append(run)
        run = OxmlElement('w:r')
        if label and any(word in label.lower() for word in ('code', 'snippet', 'expression', 'rdl text')):
            prop = OxmlElement('w:rPr'); font = OxmlElement('w:rFonts'); font.set(qn('w:ascii'), 'Consolas'); font.set(qn('w:hAnsi'), 'Consolas'); prop.append(font); run.append(prop)
        for i, line in enumerate(clean(value).split('\n')):
            if i: run.append(OxmlElement('w:br'))
            reference = reference_pattern.search(line) if links else None
            if reference:
                prefix = reference_prefixes[reference[1]]
                if prefix+reference[2] not in link_targets:
                    reference = None
            if reference:
                t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = line[:reference.start()]; run.append(t); p.append(run)
                prefix = reference_prefixes[reference[1]]
                link = OxmlElement('w:hyperlink'); link.set(qn('w:anchor'), prefix + reference[2])
                link_run = OxmlElement('w:r'); prop = OxmlElement('w:rPr')
                color = OxmlElement('w:color'); color.set(qn('w:val'), '164C8C'); prop.append(color); link_run.append(prop)
                t = OxmlElement('w:t'); t.text = reference[0]; link_run.append(t); link.append(link_run); p.append(link)
                run = OxmlElement('w:r'); line = line[reference.end():]
            t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = line; run.append(t)
        p.append(run); body.append(p)
    def lines(value, prefix='', indent=0):
        if isinstance(value, dict):
            for key, child in value.items():
                label = {'developer_note': 'Recommendations', 'remediation': 'Recommendations',
                         'recommendation': 'Recommendations', 'reviewer_note': 'Review guidance',
                         'impact': 'Security impact'}.get(key, str(key).replace('_', ' ').title())
                label = (prefix + ' / ' if prefix else '') + label
                if isinstance(child, (dict, list)):
                    yield '  '*indent + label + ':'
                    yield from lines(child, indent=indent+1)
                else:
                    yield '  '*indent + label + ': ' + clean(child if child is not None else '(not recorded)')
            if not value:
                yield '  '*indent + '(none recorded)'
        elif isinstance(value, list):
            if not value:
                yield '  '*indent + '(none recorded)'
            for index, child in enumerate(value, 1):
                yield '  '*indent + f'[{index}]'
                yield from lines(child, indent=indent+1)
        else:
            yield '  '*indent + clean(value if value is not None else '(not recorded)')
    def fields(value, prefix=''):
        # Bound paragraph size, not evidence size. Thousands of snippets can
        # otherwise produce millions of XML paragraphs and excessive memory.
        buffer, size = [], 0
        for line in lines(value, prefix):
            for offset in range(0, max(1, len(line)), 8000):
                fragment = line[offset:offset+8000]
                if size + len(fragment) > 8000 and buffer:
                    paragraph('', '\n'.join(buffer)); buffer, size = [], 0
                buffer.append(fragment); size += len(fragment)
        if buffer:
            paragraph('', '\n'.join(buffer))
    def record_table(headers, rows, widths=None):
        widths = widths or [9700//len(headers)] * len(headers)
        table = OxmlElement('w:tbl')
        properties = OxmlElement('w:tblPr'); style = OxmlElement('w:tblStyle')
        style.set(qn('w:val'), 'LightShading-Accent1'); properties.append(style)
        layout = OxmlElement('w:tblLayout'); layout.set(qn('w:type'), 'fixed'); properties.append(layout)
        table.append(properties)
        grid = OxmlElement('w:tblGrid')
        for width in widths:
            column = OxmlElement('w:gridCol'); column.set(qn('w:w'), str(width)); grid.append(column)
        table.append(grid)
        for index, values in enumerate([headers, *rows]):
            row = OxmlElement('w:tr')
            if index == 0:
                properties = OxmlElement('w:trPr'); properties.append(OxmlElement('w:tblHeader')); row.append(properties)
            for column_index, value in enumerate(values):
                cell = OxmlElement('w:tc'); properties = OxmlElement('w:tcPr')
                width = OxmlElement('w:tcW'); width.set(qn('w:w'), str(widths[column_index])); width.set(qn('w:type'),'dxa'); properties.append(width); cell.append(properties)
                p = OxmlElement('w:p')
                for offset, line in enumerate(clean(value).split('\n')):
                    if offset:
                        run = OxmlElement('w:r'); run.append(OxmlElement('w:br')); p.append(run)
                    for fragment in re.split(r'((?:'+reference_choices+r') \d+)', line):
                        reference = reference_pattern.fullmatch(fragment)
                        run = OxmlElement('w:r'); text = OxmlElement('w:t')
                        text.set(qn('xml:space'),'preserve'); text.text = fragment; run.append(text)
                        if reference:
                            prefix = reference_prefixes[reference[1]]
                            if prefix+reference[2] in link_targets:
                                properties = OxmlElement('w:rPr'); color = OxmlElement('w:color')
                                color.set(qn('w:val'), '164C8C'); properties.append(color); run.insert(0, properties)
                                link = OxmlElement('w:hyperlink'); link.set(qn('w:anchor'), prefix+reference[2]); link.append(run); p.append(link)
                            else:
                                p.append(run)
                        else:
                            p.append(run)
                cell.append(p); row.append(cell)
            table.append(row)
        body.append(table)
    sections = dict(document.get('sections', {}))
    if 'suppressed_findings' in sections:
        sections['suppressed_findings'] = context['suppressed']
    from core.rdl.narratives import enrich_source_findings
    if isinstance(sections.get('areas_of_interest'), list):
        sections['areas_of_interest'] = enrich_source_findings(sections['areas_of_interest'])
    if isinstance(sections.get('filepaths_aoi'), list):
        sections['filepaths_aoi'] = enrich_source_findings(sections['filepaths_aoi'], path_only=True)
    from utils.report_audience import filter_guidance
    sections = filter_guidance(sections, document.get('report_audience', 'both'))
    native_references = {}
    def select(row, keys):
        return {key: row[key] for key in keys if key in row}
    def flow_record(row):
        result = select(row, ('id', 'title', 'description', 'analysis_kind', 'report_platform',
                            'severity', 'cwe', 'source', 'sink', 'file', 'line', 'confidence_score',
                            'confidence_label', 'risk_score', 'input_surface', 'termination_nodes'))
        result['path'] = [select(step, ('role', 'file', 'path', 'line', 'code', 'context_before', 'context_after', 'reason'))
                          for step in row.get('path', [])]
        result['xref'] = []
        import json
        for ref in row.get('xref', []):
            reference = select(ref, ('type', 'kind', 'resolved_name', 'symbol', 'file', 'path', 'line', 'context', 'code'))
            ident = sha256(json.dumps(reference, sort_keys=True).encode()).digest()
            if ident not in native_references:
                native_references[ident] = dict(reference, reference=f'Flow reference {len(native_references)+1}')
            result['xref'].append(native_references[ident]['reference'])
        return result
    def rdl_record(row):
        result = select(row, ('title', 'rule_title', 'severity', 'location', 'report_path', 'report_kind',
                             'message', 'description', 'report_explanation', 'cwe', 'confidence',
                             'confidence_reason', 'impact', 'reviewer_note', 'remediation', 'review_area', 'source', 'sink',
                             'limitations', 'report_locations'))
        result['trace'] = select(row.get('trace') or {}, ('summary', 'review_note', 'protection_note'))
        result['review_evidence'] = []
        for step in row.get('report_evidence', []):
            evidence = select(step, ('predicate', 'location', 'matched', 'explanation', 'symbol'))
            if step.get('report_excerpt'):
                evidence['saved_source_file'] = (step.get('location') or {}).get('path') or row.get('report_path', '')
            result['review_evidence'].append(evidence)
        return result
    def semantic_step(step):
        result = select(step, ('predicate', 'location', 'symbol', 'explanation', 'transition', 'matched'))
        if step.get('report_excerpt'):
            result['saved_source_file'] = (step.get('location') or {}).get('path', '')
        return result
    def semantic_flow(row):
        result = select(row, ('title', 'rule_title', 'severity'))
        result['trace_status'] = 'Source and sink recorded' if row.get('report_complete') else 'Partial trace: destination unresolved'
        result['related_finding'] = 'RDL finding ' + row['finding_anchor'].rsplit('-', 1)[-1]
        result['steps'] = [semantic_step(step) for step in row['report_steps']]
        return result
    if isinstance(sections.get('analysis'), dict):
        sections['analysis'] = {
            **{key: value for key, value in sections['analysis'].items() if key != 'results'},
            'platform_coverage': [{key: value for key, value in target.items() if key not in ('findings', 'security_inventory')}
                                  for target in context['analysis_targets']],
            'vulnerability_assessments': context['vulnerabilities'], 'mitigation_observations': context['mitigations'],
            'manual_reviews': context['manual_reviews'], 'attack_vectors': context['attack_vectors'],
            'taint_and_data_flows': [flow_record(row) for row in context['analysis_flows']],
            'rdl_semantic_inputs': [select(row, ('predicate', 'location', 'symbol', 'explanation', 'snippet', 'finding_title'))
                                    for row in context['semantic_vectors']],
            'rdl_semantic_paths': [semantic_flow(row) for row in context['semantic_flows']],
            'rdl_call_handoffs': [dict(before=semantic_step(row['before']), after=semantic_step(row['after']),
                                      transition=row['transition'], related_flow='RDL flow '+row['flow_anchor'].rsplit('-', 1)[-1])
                                  for row in context['semantic_references']]}
    # Print review findings and coverage outcomes rather than thousands of
    # duplicated nonmatching predicate records. The HTML package retains raw JSON.
    if 'rdl_review_evidence' in sections:
        sections['rdl_review_evidence'] = {'outcomes': context['rdl_outcomes'],
            'distinct_findings': context['rdl_finding_count'],
            'findings': [rdl_record(row) for row in context['rdl_findings']],
            'coverage_limits': context['review_limitations'], 'saved_source_evidence': context['evidence_sources']}
    if 'rdl_cross_references' in sections:
        # Match the HTML/PDF's human-readable index rather than printing raw
        # graph serialization (including duplicate snippets) for every call.
        graph = sections['rdl_cross_references']
        sections['rdl_cross_references'] = {
            'summary': graph.get('summary', {}),
            'truncated': bool(graph.get('truncated')),
            'symbols': [{key: row.get(key) for key in ('name','qualified_name','kind','path','line','signature','callers','callees')}
                        for row in graph.get('symbols', [])],
            'calls': [{key: row.get(key) for key in ('callee','path','line','arguments','resolution','reason')}
                      for row in graph.get('calls', [])]}
    if native_references:
        sections['native_reference_catalogue'] = list(native_references.values())
        contents_link('native_reference_catalogue', labels['native_reference_catalogue'])
    # Catalogue source text per file, including overlapping context windows.
    # Deduplicating whole windows still repeats their individual source lines
    # and can exceed word processors' text limits on a large saved scan.
    source_keys = {'snippet', 'context', 'context_before', 'context_after', 'code', 'text'}
    excerpts = {}
    review_texts = {}
    review_counts = Counter()
    def count_reviews(value, key=None):
        if isinstance(value, str):
            if len(value) >= 200 and key not in source_keys and key not in ('title', 'rule_title'):
                review_counts[value] += 1
        elif isinstance(value, dict):
            for name, child in value.items():
                count_reviews(child, name)
        elif isinstance(value, list):
            for child in value:
                count_reviews(child, key)
    count_reviews(sections)
    def review_text(value):
        if not isinstance(value, str) or (len(value) < 8000 and review_counts[value] < 2):
            return value
        ident = sha256(value.encode()).digest()
        if ident not in review_texts:
            review_texts[ident] = {'reference': f'Review text {len(review_texts)+1}', 'original_text': value}
        return review_texts[ident]['reference']
    def walk(value, path='', line=None):
        if isinstance(value, dict):
            location = value.get('location') if isinstance(value.get('location'), dict) else {}
            candidate = value.get('file') or value.get('path')
            path = str((candidate if isinstance(candidate, str) else None) or
                       location.get('path') or location.get('file') or value.get('report_path') or path)
            line = value.get('line') or location.get('line') or line
            result = {}
            for key, child in value.items():
                if key in source_keys and isinstance(child, str) and len(child) >= 40:
                    if path not in excerpts:
                        excerpts[path] = {'reference': f'Source excerpt {len(excerpts)+1}',
                                          'path': path or 'Path not recorded', 'blocks': {}}
                    entry = excerpts[path]
                    ident = sha256(child.encode()).digest()
                    block = entry['blocks'].setdefault(ident, {'text': child, 'lines': set()})
                    if line is not None:
                        block['lines'].add(str(line))
                    result[key] = entry['reference']
                elif isinstance(child, str) and key not in ('title', 'rule_title'):
                    result[key] = review_text(child)
                else:
                    result[key] = walk(child, path, line)
            return result
        if isinstance(value, list):
            return [walk(child, path, line) for child in value]
        if isinstance(value, str):
            return review_text(value)
        return value
    sections = walk(sections)
    def bind_sources(value):
        if isinstance(value, dict):
            path = value.get('saved_source_file')
            if path in excerpts:
                value['saved_source'] = excerpts[path]['reference'] + ' · ' + path
                del value['saved_source_file']
            for child in value.values():
                bind_sources(child)
        elif isinstance(value, list):
            for child in value:
                bind_sources(child)
    bind_sources(sections)
    if review_texts:
        sections['review_text_catalogue'] = list(review_texts.values())
        contents_link('review_text_catalogue', labels['review_text_catalogue'])
    if excerpts:
        sections['source_excerpt_catalogue'] = list(excerpts.values())
        contents_link('source_excerpt_catalogue', labels['source_excerpt_catalogue'])
    for prefix, records in [('source_excerpt_', excerpts), ('flow_reference_', native_references), ('review_text_', review_texts)]:
        link_targets.update(prefix+str(i) for i in range(1, len(records)+1))
    link_targets.update('rdl_finding_'+str(i) for i in range(1, len(context['rdl_findings'])+1))
    link_targets.update('rdl_flow_'+str(i) for i in range(1, len(context['semantic_flows'])+1))
    bookmark_number = 10000
    for section_index, (key, value) in enumerate(sections.items(), 1):
        report_progress(progress, 'building', 'Building ' + labels.get(key, key.replace('_',' ').title()), section_index-1, len(sections), 'sections')
        heading(labels.get(key, key.replace('_',' ').title()), 1, (section_index, 'section_' + key))
        if key == 'program_graphs':
            from core.rdl.program_graphs import report_views
            paragraph('Snapshot', str(value.get('provenance','Not recorded'))+' · '+str(value.get('created_at','')))
            paragraph('', 'CFG records modeled control transfers; DFG records value dependencies. These are not vulnerability verdicts. Full node and edge tables are included. Use the HTML explorer or SVG export for an interactive graphical view.')
            if value.get('truncated'):
                paragraph('Coverage', 'The original graph capture reached its recorded analysis budget.')
            for note in value.get('diagnostics', []):
                paragraph('Limitation', note)
            record_table(['Language / format','CFG','DFG','Scope'], [[row['language'],row['cfg'],row['dfg'],row['note']] for row in value.get('coverage',[])])
            for view in report_views(value):
                heading(view['kind'].upper()+' · '+view['path'], 2)
                paragraph('Model status', view.get('status','Not recorded'))
                for note in view.get('diagnostics', []) + view.get('limitations', []):
                    paragraph('Limitation', note)
                record_table(['Node / kind','Operation','Source / code'], [
                    [row['id']+'\n'+row.get('kind',''),row.get('label',''),
                     str(row.get('location',{}).get('path',''))+':'+str(row.get('location',{}).get('line',''))+'\n'+row.get('code','')]
                    for row in view.get('nodes',[])])
                record_table(['From node','Relationship','To node'], [[edge['source'],edge.get('label') or edge['kind'],edge['target']] for edge in view.get('edges',[])])
            continue
        if key == 'recon':
            recon = context['recon_report']
            number = lambda value: str(value) if value is not None else 'Not recorded'
            paragraph('', 'Detected technology indicators from saved results. Confidence describes detection evidence, not vulnerability severity. Files and directories may support several technologies; match totals include those overlaps.')
            record_table(['Categories', 'Technology detections', 'Recorded file matches', 'Recorded directory matches'],
                         [[str(recon['total_categories']), str(recon['total_technologies']), str(recon['file_observations']), str(recon['directory_observations'])]])
            for item in recon['metadata']:
                paragraph(item['label'], item['value'])
            heading('Technology overview', 2)
            if not recon['technologies']:
                paragraph('', 'No technology detections were saved in this reconnaissance result.')
            for offset in range(0, len(recon['technologies']), 100):
                rows = []
                for tech in recon['technologies'][offset:offset+100]:
                    breakdown = '\n'.join(level.capitalize()+': '+number(tech['confidence_counts'][level]) for level in ('high','medium','low'))
                    rows.append([tech['category'], tech['name'], number(tech['files']), number(tech['directory_count']),
                                 tech['confidence'].capitalize()+'\n'+breakdown])
                record_table(['Category', 'Technology', 'Files', 'Directories', 'Confidence / file breakdown'], rows,
                             [1900, 1900, 900, 1100, 3900])
            paragraph('', 'Evidence paths inside the target are relative to the target directory above. Other paths retain their original form. Sample files are saved examples, not a complete inventory.')
            for category in recon['categories']:
                heading(category['name'], 2)
                for tech in category['technologies']:
                    heading(tech['name'], 3)
                    paragraph('Detection', f"{number(tech['files'])} files; {number(tech['directory_count'])} directories; {tech['confidence']} confidence")
                    record_table(['High confidence files', 'Medium confidence files', 'Low confidence files'],
                                 [[number(tech['confidence_counts'][level]) for level in ('high','medium','low')]])
                    paragraph('Directory evidence', f"{len(tech['directories'])} saved entries")
                    if not tech['directories']:
                        paragraph('', 'No directory evidence was saved for this detection.')
                    for offset in range(0, len(tech['directories']), 200):
                        record_table(['Directory', 'Matched files'],
                                     [[row['display'], number(row['files'])] for row in tech['directories'][offset:offset+200]], [7800, 1900])
                    paragraph(tech['evidence_label'], f"{len(tech['samples'])} saved entries")
                    if not tech['samples']:
                        paragraph('', 'No sample file paths were saved for this detection.')
                    for row in tech['samples']:
                        paragraph('File', row['display'], links=False)
                        if row['confidence'] != 'not recorded':
                            paragraph('Recorded file confidence', row['confidence'].capitalize())
            continue
        if key == 'rdl_cross_references':
            paragraph('', 'Saved static call graph. Unresolved calls and zero callers retain their recorded limitations. Complete original graph metadata is available in the JSON export and HTML package.')
            if value['truncated']:
                paragraph('Coverage limit', 'The saved cross-reference graph reports truncation. This report preserves the available saved entries.')
            fields(value['summary'])
            for field, label in [('symbols','Symbol Index'), ('calls','Call-site Index')]:
                records = value[field]
                for offset in range(0, len(records), 200):
                    heading(f'{label}: {offset+1}-{min(offset+200,len(records))} of {len(records)}', 2)
                    rows = []
                    for row in records[offset:offset+200]:
                        location = str(row.get('path') or '-') + (':'+str(row['line']) if row.get('line') else '')
                        if field == 'symbols':
                            rows.append([str(row.get('qualified_name') or row.get('name') or '')+'\n'+str(row.get('kind') or ''),
                                         location+'\n'+str(row.get('signature') or ''),
                                         f"{row.get('callers') or 0} / {row.get('callees') or 0}"])
                        else:
                            arguments = [str(argument.get('name') or argument.get('parameter') or 'Argument') + ': ' +
                                         str(argument.get('expression') or '') for argument in row.get('arguments') or []]
                            rows.append([row.get('callee') or '',location + ('\n'+'\n'.join(arguments) if arguments else ''),
                                         str(row.get('resolution') or 'Not recorded')+'\n'+str(row.get('reason') or '')])
                    headers = ['Symbol / kind','Definition / signature','Callers / callees'] if field=='symbols' else ['Callee','Call site','Resolution / reason']
                    record_table(headers, rows)
            continue
        if isinstance(value, list):
            paragraph('', f'{len(value):,} saved records. All recorded evidence is included.')
            for i, row in enumerate(value, 1):
                if i % 200 == 0:
                    report_progress(progress, 'building', 'Building ' + labels.get(key, key.replace('_',' ').title()), i, len(value), 'records')
                title = next((str(row[k]) for k in ('reference','rule_title','title','name','v2_rule_id') if row.get(k)), f'Record {i}') if isinstance(row,dict) else f'Record {i}'
                catalogue_prefix = {'source_excerpt_catalogue':'source_excerpt_',
                                    'native_reference_catalogue':'flow_reference_',
                                    'review_text_catalogue':'review_text_'}.get(key)
                bookmark = (bookmark_number, catalogue_prefix+str(i)) if catalogue_prefix else None
                if bookmark:
                    bookmark_number += 1
                heading(f'{i}. {title}', 2, bookmark)
                if key == 'source_excerpt_catalogue':
                    paragraph('File', row['path'])
                    for block in row['blocks'].values():
                        positions = ', '.join(sorted(block['lines'], key=lambda v: (not v.isdigit(), int(v) if v.isdigit() else v)))
                        paragraph('Recorded lines', positions or 'Not recorded')
                        # Bound paragraph length while preserving every character.
                        for offset in range(0, len(block['text']), 8000):
                            paragraph('Code' if offset == 0 else 'Code continued', block['text'][offset:offset+8000], links=False)
                elif key == 'review_text_catalogue':
                    for offset in range(0, len(row['original_text']), 8000):
                        paragraph('Original text' if offset == 0 else 'Continued text', row['original_text'][offset:offset+8000], links=False)
                else:
                    fields(row)
        elif isinstance(value, dict):
            for name, records in value.items():
                heading(str(name).replace('_', ' ').title(), 2)
                if isinstance(records, list):
                    paragraph('', f'{len(records):,} saved entries')
                    for index, row in enumerate(records, 1):
                        if index % 200 == 0:
                            report_progress(progress, 'building', 'Building ' + str(name).replace('_', ' '), index, len(records), 'records')
                        prefix = ('rdl_finding_' if key == 'rdl_review_evidence' and name == 'findings' else
                                  'rdl_flow_' if key == 'analysis' and name == 'rdl_semantic_paths' else None)
                        if prefix:
                            title = row.get('title') or row.get('rule_title') or f'Record {index}'
                            heading(f'{index}. {title}', 3, (bookmark_number, prefix+str(index)))
                            bookmark_number += 1
                        fields(row, f'{index}')
                else:
                    fields(records)
        else:
            fields(value)
    body.append(section_properties)
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    report_progress(progress, 'packaging', 'Packaging the Word document')
    doc.save(destination)
    return destination
